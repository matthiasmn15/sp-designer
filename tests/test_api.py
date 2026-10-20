import io

from openpyxl import load_workbook


def test_ui_and_assets_served(client):
    assert "Spark Plug Designer" in client.get("/").text
    assert client.get("/static/app.js").status_code == 200
    assert client.get("/vendor/plotly.min.js").status_code == 200


def test_materials_endpoints(client):
    mats = client.get("/api/materials", params={"category": "precious"}).json()
    ids = {m["id"] for m in mats}
    assert {"Ir", "Pt", "Ru", "Rh", "Ir-10Rh"} <= ids
    assert "Alloy-600" not in ids
    assert client.get("/api/materials/nope").status_code == 404


def test_override_and_reset(client):
    r = client.put("/api/materials/Pt/properties", json={"bead_tendency": {"value": 2, "source": "test"}})
    assert r.status_code == 200
    assert r.json()["properties"]["bead_tendency"]["origin"] == "user"
    assert client.put("/api/materials/Pt/properties", json={"price": {"value": 1}}).status_code == 400
    r = client.delete("/api/materials/Pt/properties")
    assert "bead_tendency" not in r.json()["properties"]


def test_add_material_percent_composition(client):
    r = client.post("/api/materials", json={"name": "Ir 5Rh", "kind": "alloy", "categories": ["precious"],
                                            "composition": {"Ir": 95, "Rh": 5}})
    assert r.status_code == 201
    body = r.json()
    assert body["id"] == "Ir-5Rh"
    assert abs(body["composition"]["Rh"] - 0.05) < 1e-9
    assert client.delete("/api/materials/Ir-5Rh").status_code == 204


def test_ranking_endpoint(client):
    r = client.post("/api/ranking", json={"weights": {"melting_point": 1}, "category": "precious"})
    assert r.json()[0]["id"] == "Ir"


def test_manual_price_updates_material(client):
    r = client.post("/api/prices", json={"metal": "Ni", "value": 0.02, "unit": "USD/g", "date": "2026-10-20"})
    assert r.status_code == 201
    ni = client.get("/api/materials/Ni").json()
    assert abs(ni["properties"]["price"]["value"] - 0.02) < 1e-9
    assert client.post("/api/prices", json={"metal": "Alloy-600", "value": 1}).status_code == 400


def test_exports(client):
    csv = client.get("/api/export", params={"ids": "Ir,Pt", "format": "csv"})
    assert csv.status_code == 200
    assert csv.text.splitlines()[0] == "Group;Property;Unit;Iridium;Platinum"
    xlsx = client.get("/api/export", params={"ids": "Ir,Pt", "format": "xlsx"})
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames == ["Comparison", "Sources"]
    assert wb["Comparison"]["D1"].value == "Iridium"


def test_settings_roundtrip(client):
    r = client.put("/api/settings", json={"price_provider": "none", "price_max_age_days": 7})
    assert r.json()["price_max_age_days"] == 7
    assert client.post("/api/prices/refresh").status_code == 400
