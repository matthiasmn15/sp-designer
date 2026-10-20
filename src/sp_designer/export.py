"""CSV and Excel export of material comparisons."""

from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .materials.models import Material
from .materials.properties import PROPERTIES


def comparison_rows(materials: list[Material]) -> list[list]:
    header = ["Group", "Property", "Unit", *[m.name for m in materials]]
    rows: list[list] = [header]
    for p in PROPERTIES:
        values = []
        for m in materials:
            pv = m.properties.get(p.key)
            values.append(None if pv is None or pv.value is None else round(pv.value, max(p.decimals, 3)))
        rows.append([p.group, p.label, p.unit, *values])
    return rows


def source_rows(materials: list[Material]) -> list[list]:
    rows: list[list] = [["Material", "Property", "Value", "Unit", "Origin", "Source", "Date", "Note"]]
    for m in materials:
        for p in PROPERTIES:
            pv = m.properties.get(p.key)
            if pv is None or pv.value is None:
                continue
            rows.append([m.name, p.label, pv.value, p.unit, pv.origin, pv.source, pv.date, pv.note])
    return rows


def to_csv(materials: list[Material]) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";")  # semicolon: opens correctly in German-locale Excel
    writer.writerows(comparison_rows(materials))
    return buf.getvalue()


def to_xlsx(materials: list[Material]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"
    _write(ws, comparison_rows(materials))
    _write(wb.create_sheet("Sources"), source_rows(materials))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write(ws, rows: list[list]) -> None:
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B4FC4")
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    for idx, column in enumerate(ws.columns, start=1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in column)
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 10), 60)
